"""Monthly report of one battery: Excel workbook (openpyxl) and a self-contained HTML page.

Contents: KPIs of the month next to the guarantees (RTE with and without auxiliaries, availability, cycles,
throughput, SoC, C-rate, temperatures), the warranty status, the warranty violations of the month, the illustrative
liquidated damages of the LTSA year, capacity tests and the SoH forecast, the energy balance and the auxiliary
model. Every contractual figure is marked as illustrative.

Usage:
    write_excel(res, "2026-08", Path("out/report_2026-08.xlsx"))
    write_html(res, "2026-08", Path("out/report_2026-08.html"))
"""
from __future__ import annotations

import base64
import html
import io
import logging
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

log = logging.getLogger(__name__)
logging.getLogger("matplotlib").setLevel(logging.WARNING)
DISCLAIMER = ("Contract terms (warranty, LTSA, liquidated damages) are illustrative placeholders, not a reading of any "
              "contract; check them against the actual documents.")
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RED = PatternFill("solid", fgColor="F4C7C3")
GREEN = PatternFill("solid", fgColor="D9EAD3")
AMBER = PatternFill("solid", fgColor="FCE8B2")
RULE_LABELS = {
    "efc_per_year": "Cycles per warranty year", "efc_per_day": "Cycles per day", "throughput_total": "Throughput",
    "soc_low": "SoC below window", "soc_high": "SoC above window", "mean_soc": "Mean SoC (resting)",
    "c_rate": "C-rate", "cell_temp_high": "Cell temperature high", "cell_temp_low": "Cell temperature low",
    "cell_temp_annual": "Annual mean cell temperature", "ambient_high": "Ambient high", "ambient_low": "Ambient low",
}


def month_slice(res: dict, month: str) -> dict:
    """The rows of every result table that belong to ``month`` (YYYY-MM), and the state at its end (warranty
    status, LTSA year to date, SoH forecast, LD) as far as it was known then."""
    from bess_ops.pipeline import as_of
    tz = res["site_cfg"]["timezone"]
    k = res["monthly"]
    if month not in k.index:
        raise ValueError(f"month {month} not in the data ({k.index[0]} .. {k.index[-1]})")
    cache = res.setdefault("_as_of", {})
    if month not in cache:
        cache[month] = as_of(res, month)
    st = cache[month]
    v = res["violations"]
    starts = pd.DatetimeIndex(v["start"]).tz_convert(tz).strftime("%Y-%m") if len(v) else pd.Index([])
    vm = v[starts == month] if len(v) else v
    ongoing = v[v["rule"].isin(["efc_per_year", "throughput_total", "cell_temp_annual"]) & (starts <= month)] \
        if len(v) else v
    days = res["daily"][res["daily"].index.str[:7] == month]
    tests = res["tests"]
    if len(tests):
        tests = tests[pd.DatetimeIndex(tests["start"]).tz_convert(tz).strftime("%Y-%m") <= month]
    return {"kpi": k.loc[month], "violations": vm, "ongoing": ongoing, "daily": days, "tests": tests,
            "state": st, "balance": res["balance"].loc[month] if month in res["balance"].index else None}


def kpi_rows(res: dict, month: str) -> list[tuple]:
    """(label, value, guarantee or limit, status, unit) for the KPI table."""
    m = month_slice(res, month)
    k = m["kpi"]
    lt, w = res["ltsa"], res["warranty"]
    rows = []

    def add(label, value, ref=None, ok=None, unit="", bad="below"):
        rows.append((label, value, ref, "" if ok is None else ("ok" if ok else bad), unit))

    add("Energy charged at the meter", k["e_in_mwh"], unit="MWh")
    add("Energy discharged at the meter", k["e_out_mwh"], unit="MWh")
    if pd.notna(k.get("e_aux_mwh")):
        add("Auxiliary energy", k["e_aux_mwh"], unit="MWh")
    g = lt["rte"]["guaranteed"]
    defin = lt["rte"]["definition"]
    rte_c = k["rte"] if defin == "incl_aux" else k["rte_excl_aux"]
    add(f"RTE, SoC-corrected ({'incl.' if defin == 'incl_aux' else 'excl.'} auxiliaries, LTSA definition)", rte_c, g,
        bool(rte_c >= g) if pd.notna(rte_c) else None)
    other = k["rte_excl_aux"] if defin == "incl_aux" else k["rte"]
    add(f"RTE, SoC-corrected ({'excl.' if defin == 'incl_aux' else 'incl.'} auxiliaries)", other)
    add("RTE, metered out / in (uncorrected)", k["rte_raw"])
    if "availability" in k and pd.notna(k["availability"]):
        ga = lt["availability"]["guaranteed"]
        add(f"Availability of the month ({lt['availability']['formula']} formula)", k["availability"], ga,
            bool(k["availability"] >= ga))
    ay = m["state"]["availability_year"]
    if ay is not None and len(ay):
        add("Availability, LTSA year to date", ay["availability"].iloc[-1], lt["availability"]["guaranteed"],
            bool(ay["availability"].iloc[-1] >= lt["availability"]["guaranteed"]))
    add("Equivalent full cycles (warranty basis)", k["efc"], unit="EFC")
    if "rainflow_efc" in k:
        add("Rainflow cycles (sum of depths)", k["rainflow_efc"], unit="EFC")
    st = m["state"]["warranty_status"]
    lim = w["throughput"]["efc_per_year"]
    add("Cycles this warranty year so far", st["efc_this_year"], lim, bool(st["efc_this_year"] <= lim) if lim
        else None, "EFC", bad="above")
    add("Cycles this warranty year, projected to a full year", st["efc_this_year_projected"], lim,
        bool(st["efc_this_year_projected"] <= lim) if lim else None, "EFC", bad="above")
    if "mean_soc" in k:
        add("Mean SoC", k["mean_soc"])
    if "hours_soc_above_max" in k:
        add("Hours above the SoC window", k["hours_soc_above_max"], unit="h")
    if "hours_soc_below_min" in k:
        add("Hours below the SoC window", k["hours_soc_below_min"], unit="h")
    cr = w["c_rate"]["max"]
    add("Highest C-rate", k["max_c_rate"], cr, bool(k["max_c_rate"] <= cr) if cr else None, "1/h", bad="above")
    if "t_cell_max_max" in k:
        tl = w["temperature"]["cell_max_c"]
        add("Highest cell temperature", k["t_cell_max_max"], tl, bool(k["t_cell_max_max"] <= tl) if tl else None,
            "degC", bad="above")
    if "t_amb_mean" in k:
        add("Mean ambient temperature", k["t_amb_mean"], unit="degC")
    add("Warranty violations started this month", int(k["violations"]), 0, bool(k["violations"] == 0), bad="above")
    add("Guaranteed SoH now (equivalent age)", st["guaranteed_soh_now"])
    return rows


def _fmt(v, unit=""):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if isinstance(v, (int, np.integer)):
        return f"{v:d}"
    if unit in ("MWh",):
        return f"{v:,.1f}"
    if unit in ("h", "EFC", "degC"):
        return f"{v:,.1f}"
    if unit == "1/h":
        return f"{v:.3f}"
    if abs(v) <= 1.5:
        return f"{v:.1%}" if unit == "" else f"{v:.3f}"
    return f"{v:,.2f}"


def _local(ts, tz):
    return pd.DatetimeIndex(ts).tz_convert(tz).strftime("%Y-%m-%d %H:%M")


def violation_table(res: dict, month: str) -> pd.DataFrame:
    m = month_slice(res, month)
    tz = res["site_cfg"]["timezone"]
    v = pd.concat([m["violations"], m["ongoing"]]).drop_duplicates(subset=["rule", "start"])
    if v.empty:
        return pd.DataFrame(columns=["Rule", "Start (local)", "End (local)", "Hours", "Peak / total", "Limit",
                                     "Overlaps logged event", "Note"])
    return pd.DataFrame({
        "Rule": v["rule"].map(RULE_LABELS).fillna(v["rule"]), "Start (local)": _local(v["start"], tz),
        "End (local)": _local(v["end"], tz), "Hours": v["duration_h"].round(2), "Peak / total": v["value"].round(3),
        "Limit": v["limit"], "Overlaps logged event": v["context"].fillna(""), "Note": v["detail"].fillna(""),
    })


def ld_table(res: dict, month: str) -> pd.DataFrame:
    ld = month_slice(res, month)["state"]["ld"]
    return ld.rename(columns={"ltsa_year": "LTSA year", "availability": "Availability LD", "rte": "RTE LD",
                              "capacity": "Capacity LD", "total": "Total", "cap": "Cap", "capped_total": "Payable",
                              "currency": "Currency"})


def test_table(res: dict, month: str | None = None) -> pd.DataFrame:
    t = month_slice(res, month)["tests"] if month else res["tests"]
    if t is None or len(t) == 0:
        return pd.DataFrame()
    tz = res["site_cfg"]["timezone"]
    return pd.DataFrame({
        "Start (local)": _local(t["start"], tz), "Duration h": t["duration_h"].round(2),
        "Cell temp degC": t["temp_c"].round(1), "Measured MWh": t["e_meas_mwh"].round(2),
        "Corrected to ref. temp MWh": t["e_corr_mwh"].round(2), "Expanded uncertainty MWh": t["U_mwh"].round(2),
        "SoH": t["soh"].round(4), "Guaranteed MWh": t["e_guaranteed_mwh"].round(2),
        "Margin MWh": t["margin_mwh"].round(2), "Verdict": t["verdict"], "Procedure checks": t["procedure_checks"],
    })


# ---------------------------------------------------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------------------------------------------------

def _sheet(wb, title, df: pd.DataFrame, widths=None, index=False):
    ws = wb.create_sheet(title)
    frame = df.reset_index() if index else df
    ws.append([str(c) for c in frame.columns])
    for row in frame.itertuples(index=False):
        ws.append([_cell(v) for v in row])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for i, col in enumerate(frame.columns, start=1):
        w = (widths or {}).get(col) or min(max(len(str(col)), *(len(str(_cell(v))) for v in frame[col].head(200)),
                                               8) + 2, 50)
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _cell(v):
    if isinstance(v, (pd.Timestamp,)):
        return v.tz_convert(None).to_pydatetime() if v.tzinfo else v.to_pydatetime()
    if isinstance(v, (np.floating,)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, float) and np.isnan(v):
        return None
    return v


def write_excel(res: dict, month: str, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI"
    ws.append([f"{res['site_cfg']['name']}: monthly report {month}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([DISCLAIMER])
    ws["A2"].font = Font(italic=True, color="7F7F7F")
    ws.append([])
    ws.append(["KPI", "Value", "Guarantee / limit", "Status", "Unit"])
    for c in ws[4]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for label, value, ref, status, unit in kpi_rows(res, month):
        ws.append([label, _cell(value), _cell(ref), status, unit])
        r = ws.max_row
        fmt = "0.0%" if unit == "" and isinstance(value, (float, np.floating)) and abs(value) <= 1.5 else "#,##0.0"
        if unit == "1/h":
            fmt = "0.000"
        for col in ("B", "C"):
            ws[f"{col}{r}"].number_format = fmt
    last = ws.max_row
    ws.conditional_formatting.add(f"D5:D{last}", CellIsRule(operator="equal", formula=['"below"'], fill=RED))
    ws.conditional_formatting.add(f"D5:D{last}", CellIsRule(operator="equal", formula=['"above"'], fill=RED))
    ws.conditional_formatting.add(f"D5:D{last}", CellIsRule(operator="equal", formula=['"ok"'], fill=GREEN))
    ws.column_dimensions["A"].width = 70
    for col in "BCDE":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "A5"

    k = res["monthly"].copy()
    ws_m = _sheet(wb, "Monthly", k.round(4), index=True)
    cols = list(k.reset_index().columns)
    n = len(k) + 1
    for name, ref_col in (("rte", "rte_guaranteed"), ("availability", "availability_guaranteed")):
        if name in cols and ref_col in cols:
            c = get_column_letter(cols.index(name) + 1)
            g = get_column_letter(cols.index(ref_col) + 1)
            ws_m.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="lessThan", formula=[f"${g}2"],
                                                                         fill=RED))
            ws_m.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="greaterThanOrEqual",
                                                                         formula=[f"${g}2"], fill=GREEN))
    if "violations" in cols:
        c = get_column_letter(cols.index("violations") + 1)
        ws_m.conditional_formatting.add(f"{c}2:{c}{n}", CellIsRule(operator="greaterThan", formula=["0"], fill=AMBER))
    for name in ("rte", "rte_excl_aux", "rte_raw", "availability", "mean_soc", "coverage"):
        if name in cols:
            c = get_column_letter(cols.index(name) + 1)
            for row in range(2, n + 1):
                ws_m[f"{c}{row}"].number_format = "0.0%"

    d = month_slice(res, month)["daily"][["e_in", "e_out", "e_aux", "d_stored", "efc", "rte_raw", "rte",
                                          "rte_excl_aux", "valid"]]
    ws_d = _sheet(wb, "Daily", d.round(4), index=True)
    c = get_column_letter(list(d.reset_index().columns).index("rte") + 1)
    ws_d.conditional_formatting.add(f"{c}2:{c}{len(d) + 1}", ColorScaleRule(start_type="min", start_color="F8696B",
                                                                           mid_type="percentile", mid_value=50,
                                                                           mid_color="FFEB84", end_type="max",
                                                                           end_color="63BE7B"))
    vt = violation_table(res, month)
    ws_v = _sheet(wb, "Violations", vt)
    if len(vt):
        ws_v.conditional_formatting.add(f"A2:H{len(vt) + 1}", CellIsRule(operator="equal", formula=['""'],
                                                                         fill=PatternFill()))
    ldt = ld_table(res, month)
    ws_l = _sheet(wb, "LD (illustrative)", ldt)
    if len(ldt):
        c = get_column_letter(list(ldt.columns).index("Payable") + 1)
        ws_l.conditional_formatting.add(f"{c}2:{c}{len(ldt) + 1}",
                                        CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
    items = month_slice(res, month)["state"]["ld"].attrs.get("items")
    if items is not None and len(items):
        _sheet(wb, "LD items", items.round(4))
    tt = test_table(res, month)
    if len(tt):
        ws_t = _sheet(wb, "Capacity tests", tt)
        col = get_column_letter(list(tt.columns).index("Verdict") + 1)
        ws_t.conditional_formatting.add(f"{col}2:{col}{len(tt) + 1}",
                                        CellIsRule(operator="equal", formula=['"fail"'], fill=RED))
        ws_t.conditional_formatting.add(f"{col}2:{col}{len(tt) + 1}",
                                        CellIsRule(operator="equal", formula=['"inconclusive"'], fill=AMBER))
        ws_t.conditional_formatting.add(f"{col}2:{col}{len(tt) + 1}",
                                        CellIsRule(operator="equal", formula=['"pass"'], fill=GREEN))
    fc = month_slice(res, month)["state"]["forecast"]
    if fc:
        _sheet(wb, "SoH forecast", fc["curve"].round(4))
    _sheet(wb, "Energy balance", res["balance"].round(3), index=True)
    if res.get("availability_month") is not None:
        _sheet(wb, "Availability", res["availability_month"].round(4), index=True)
    if res.get("aux"):
        p = res["aux"]["params"]
        _sheet(wb, "Auxiliaries", pd.DataFrame({"parameter": list(p), "value": [_cell(v) for v in p.values()]}))
    q = res["quality"].reset_index()
    _sheet(wb, "Data quality", q)
    wb.save(path)
    log.info("Excel report %s", path)
    return path


# ---------------------------------------------------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------------------------------------------------

def _png(fig) -> str:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=110, bbox_inches="tight")
    plt.close(fig)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


def _table(df: pd.DataFrame, cls="") -> str:
    if df is None or len(df) == 0:
        return "<p class='muted'>None.</p>"
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in df.columns)
    body = []
    for row in df.itertuples(index=False):
        cells = []
        for v in row:
            s = "" if v is None or (isinstance(v, float) and np.isnan(v)) else (
                f"{v:,.3f}" if isinstance(v, (float, np.floating)) else str(v))
            cells.append(f"<td>{html.escape(s)}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return f"<div class='scroll'><table class='{cls}'><thead><tr>{head}</tr></thead><tbody>{''.join(body)}" \
           f"</tbody></table></div>"


def _charts(res: dict, month: str) -> list[tuple[str, str]]:
    out = []
    d = month_slice(res, month)["daily"]
    lt = res["ltsa"]
    fig, ax = plt.subplots(figsize=(8, 3))
    x = pd.to_datetime(d.index)
    ax.bar(x, d["e_out"], color="#2a78d6", label="discharged (MWh)")
    ax.bar(x, -d["e_in"], color="#eb6834", label="charged (MWh)")
    ax2 = ax.twinx()
    ax2.plot(x, d["rte"].where(d["valid"]), "k.-", label="RTE incl. aux (valid days)")
    ax2.axhline(lt["rte"]["guaranteed"], color="#c8423b", ls="--", lw=1, label="guaranteed (monthly)")
    ax2.set_ylim(0.6, 1.0)
    ax.set_ylabel("MWh")
    ax2.set_ylabel("RTE")
    ax.legend(loc="lower left", fontsize=7)
    ax2.legend(loc="lower right", fontsize=7)
    ax.set_title(f"Daily energy and RTE, {month}", fontsize=10)
    fig.autofmt_xdate()
    out.append(("Daily energy and RTE", _png(fig)))
    k = res["monthly"]
    fig, ax = plt.subplots(figsize=(8, 3))
    ax.plot(k.index, k["rte"], "o-", color="#2a78d6", label="RTE incl. aux")
    ax.plot(k.index, k["rte_excl_aux"], "o-", color="#1baf7a", label="RTE excl. aux")
    if "availability" in k:
        ax.plot(k.index, k["availability"], "s-", color="#8e7cc3", label="availability")
    ax.axhline(lt["rte"]["guaranteed"], color="#c8423b", ls="--", lw=1)
    ax.axhline(lt["availability"]["guaranteed"], color="#8e7cc3", ls=":", lw=1)
    ax.axvline(month, color="#999", lw=6, alpha=0.3)
    ax.set_ylim(0.75, 1.01)
    ax.legend(fontsize=7, ncol=3)
    ax.tick_params(axis="x", rotation=60, labelsize=7)
    ax.set_title("Monthly RTE and availability against the guarantees (dashed)", fontsize=10)
    out.append(("Monthly trend", _png(fig)))
    state = month_slice(res, month)["state"]
    if state["forecast"]:
        fc = state["forecast"]["curve"]
        t = month_slice(res, month)["tests"]
        fig, ax = plt.subplots(figsize=(8, 3))
        ax.fill_between(fc["years"], fc["soh_p10"], fc["soh_p90"], color="#2a78d6", alpha=0.2, label="P10-P90")
        ax.plot(fc["years"], fc["soh_p50"], color="#2a78d6", label="forecast")
        ax.plot(fc["years"], fc["soh_guaranteed"], color="#c8423b", ls="--", label="guaranteed")
        req = lt["capacity"]["forecast"]["required_mwh"]
        if req:
            ax.axhline(req / res["warranty"]["reference_energy_mwh"], color="#e3a33b", ls=":", label="requirement")
        ax.errorbar(t["elapsed_years"], t["soh"], yerr=t["U_mwh"] / res["warranty"]["reference_energy_mwh"], fmt="ko",
                    ms=4, label="tests (U, k=2)")
        ax.set_xlabel("years since COD")
        ax.set_ylabel("energy / reference")
        ax.set_xlim(0, min(fc["years"].max(), 15))
        ax.set_ylim(0.75, 1.08)
        ax.legend(fontsize=7, ncol=3)
        ax.set_title("Capacity tests, SoH forecast and guaranteed curve", fontsize=10)
        out.append(("SoH", _png(fig)))
    return out


CSS = """
:root{color-scheme:light;--page:#f9f9f7;--surface:#fff;--ink:#111;--ink2:#555;--ring:#e2e1dc;--bad:#f4c7c3;
--good:#d9eAD3;--warn:#fce8b2}
@media (prefers-color-scheme:dark){:root{color-scheme:dark;--page:#111;--surface:#1b1b1a;--ink:#eee;--ink2:#bbb;
--ring:#333;--bad:#5c2a27;--good:#23412a;--warn:#4d4220}}
body{margin:0;background:var(--page);color:var(--ink);font:14px/1.45 system-ui,-apple-system,"Segoe UI",sans-serif}
main{max-width:1100px;margin:0 auto;padding:20px 16px 48px}h1{font-size:1.5rem;margin:0 0 4px}
h2{font-size:1.1rem;margin:28px 0 6px}.muted{color:var(--ink2)}.note{color:var(--ink2);font-size:.85rem}
.scroll{overflow-x:auto}table{border-collapse:collapse;width:100%;background:var(--surface);font-size:.85rem}
th,td{border:1px solid var(--ring);padding:4px 6px;text-align:left;vertical-align:top}th{background:#1f3a5f;color:#fff}
td.ok{background:var(--good)}td.below{background:var(--bad)}img{max-width:100%;height:auto;background:#fff;
border-radius:6px}.tiles{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:10px;margin:14px 0}
.tile{background:var(--surface);border:1px solid var(--ring);border-radius:8px;padding:10px}.tile .v{font-size:1.4rem;
font-weight:650}.tile .l{color:var(--ink2);font-size:.8rem}
"""


def write_html(res: dict, month: str, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    rows = kpi_rows(res, month)
    tr = []
    for label, value, ref, status, unit in rows:
        tr.append(f"<tr><td>{html.escape(label)}</td><td>{_fmt(value, unit)} {html.escape(unit)}</td>"
                  f"<td>{_fmt(ref, unit) if ref is not None else ''}</td><td class='{status}'>{status}</td></tr>")
    kpi = ("<div class='scroll'><table><thead><tr><th>KPI</th><th>Value</th><th>Guarantee / limit</th>"
           f"<th>Status</th></tr></thead><tbody>{''.join(tr)}</tbody></table></div>")
    k = res["monthly"].loc[month]
    state = month_slice(res, month)["state"]
    ld = state["ld"]
    cur = ld["currency"].iloc[0] if len(ld) else ""
    tiles = [("RTE incl. aux", f"{k['rte']:.1%}" if pd.notna(k["rte"]) else "n/a"),
             ("RTE excl. aux", f"{k['rte_excl_aux']:.1%}" if pd.notna(k["rte_excl_aux"]) else "n/a"),
             ("Availability", f"{k['availability']:.2%}" if "availability" in k and pd.notna(k["availability"])
              else "n/a"),
             ("Cycles", f"{k['efc']:.1f} EFC"), ("Violations", f"{int(k['violations'])}"),
             ("LD up to this month (illustrative)", f"{ld['capped_total'].sum():,.0f} {cur}" if len(ld) else "0")]
    tiles_html = "".join(f"<div class='tile'><div class='v'>{html.escape(v)}</div><div class='l'>{html.escape(t)}"
                         f"</div></div>" for t, v in tiles)
    charts = "".join(f"<h2>{html.escape(t)}</h2><img alt='{html.escape(t)}' src='{src}'>" for t, src in
                     _charts(res, month))
    st = state["warranty_status"]
    status = pd.DataFrame({"Item": ["Warranty year", "Cycles this warranty year", "Projected cycles this year",
                                    "Cycles since COD", "Throughput MWh", "Equivalent age (years)",
                                    "Guaranteed SoH now"],
                           "Value": [str(st["warranty_year"]), f"{st['efc_this_year']:.1f}",
                                     f"{st['efc_this_year_projected']:.1f}", f"{st['efc_total']:.1f}",
                                     f"{st['throughput_mwh']:,.0f}", f"{st['equivalent_age_years']:.2f}",
                                     f"{st['guaranteed_soh_now']:.1%}"]})
    fc_html = ""
    if state.get("handoff"):
        h = state["handoff"]
        fc_html = ("<p>Fade fitted to the tests: calendar {:.4f} per sqrt(year), cycling {:.4f} per 1000 EFC "
                   "(prior from the LFP cell fit of bess-capacity-augmentation). Below the requirement: {} "
                   "(range {} to {}); below the guaranteed curve: {}.</p>").format(
            h["fade_calendar_per_sqrt_year"], h["fade_per_1000_efc"],
            h["below_requirement"]["central"] or "not within the horizon", h["below_requirement"]["early"] or "-",
            h["below_requirement"]["late"] or "beyond the horizon", h["below_guarantee"]["central"] or "not within "
            "the horizon")
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>BESS monthly report {month}</title><style>{CSS}</style></head><body><main>
<h1>{html.escape(res['site_cfg']['name'])}: monthly report {month}</h1>
<p class="note">{html.escape(DISCLAIMER)} Times are local ({html.escape(res['site_cfg']['timezone'])}).</p>
<div class="tiles">{tiles_html}</div>
<h2>KPIs and guarantees</h2>{kpi}
<h2>Warranty status</h2>{_table(status)}
<h2>Warranty violations (started this month, and open yearly limits)</h2>{_table(violation_table(res, month))}
<p class="note">Violations that overlap a logged capacity test or outage are listed with that event; whether they
are excused depends on the warranty.</p>
<h2>Liquidated damages per LTSA year, up to the end of this month (illustrative)</h2>{_table(ld_table(res, month))}
<h2>Capacity tests up to this month</h2>{_table(test_table(res, month))}{fc_html}
{charts}
<h2>Energy balance of the month (MWh)</h2>{_table(res['balance'].loc[[month]].round(2).reset_index())}
<h2>Data quality</h2>{_table(res['quality'].round(3).reset_index())}
</main></body></html>"""
    path.write_text(page, encoding="utf-8")
    log.info("HTML report %s", path)
    return path
