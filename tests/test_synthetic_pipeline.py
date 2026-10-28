"""End to end on the synthetic battery: every estimator against the known truth, the report and the CLI."""
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from bess_ops.cli import main
from bess_ops.report import write_excel, write_html
from bess_ops.study import true_availability, true_losses
from bess_ops.synthetic import Plant


def test_rte_with_and_without_auxiliaries(demo):
    sim, frame, res = demo
    truth = sim.truth["rte_excl_aux"]
    assert res["yearly_rte"]["rte_excl_aux"].to_numpy() == pytest.approx(truth, abs=2e-4)
    assert (res["monthly_rte"]["rte_excl_aux"] - truth).abs().max() < 1e-3
    d = res["daily"][res["daily"]["valid"]]
    assert d["rte_excl_aux"].median() == pytest.approx(truth, abs=5e-4)
    raw = d["e_out_pcs"] / d["e_in_pcs"]
    # the SoC correction narrows the daily spread several times
    assert (d["rte_excl_aux"].quantile(0.95) - d["rte_excl_aux"].quantile(0.05)) < \
        0.3 * (raw.quantile(0.95) - raw.quantile(0.05))
    # winter heating pulls the RTE including auxiliaries below the illustrative 85 % guarantee
    m = res["monthly_rte"]["rte"]
    assert m.loc["2025-01"] < 0.85 < m.loc["2025-07"]


def test_auxiliaries_availability_and_tests(demo):
    sim, frame, res = demo
    p = res["aux"]["params"]
    for k, v in sim.truth["aux"].items():
        assert p[k] == pytest.approx(v, rel=0.02, abs=0.05 if k.startswith("t_") else 1e-3), k
    ay = res["availability_year"]["availability"]
    assert ay.iloc[0] == pytest.approx(true_availability(sim.truth, res["ltsa"], 2025), abs=1e-9)
    assert ay.iloc[1] == pytest.approx(true_availability(sim.truth, res["ltsa"], 2026), abs=1e-9)
    assert ay.iloc[1] < 0.97 <= ay.iloc[0]                       # the ten-day transformer outage costs year 2
    tests = res["tests"]
    assert len(tests) == 4
    for i, t in enumerate(sim.truth["tests"]):
        assert tests.iloc[i]["e_corr_mwh"] == pytest.approx(t["e_meter_25c_mwh"], abs=0.1)
        assert abs(tests.iloc[i]["e_corr_mwh"] - t["e_meter_25c_mwh"]) < tests.iloc[i]["U_mwh"]


def test_violations_found(demo):
    sim, frame, res = demo
    v = res["violations"]
    counts = v["rule"].value_counts()
    assert counts["efc_per_day"] == 14
    assert counts["c_rate"] == 1 and counts["efc_per_year"] == 1
    assert counts["cell_temp_high"] >= 1
    tz = "Europe/Warsaw"
    misconfig = v[(v["rule"] == "soc_high") & (v["context"] == "")]
    assert len(misconfig) == 1
    assert misconfig["start"].iloc[0].tz_convert(tz).strftime("%Y-%m-%d") == "2026-05-01"
    assert (v.loc[v["rule"] == "soc_high", "context"] == "capacity_test").sum() == 4
    assert set(counts.index) == {"efc_per_day", "c_rate", "efc_per_year", "cell_temp_high", "soc_high"}


def test_energy_balance_closes(demo):
    sim, frame, res = demo
    tl = true_losses(sim.frame, Plant(), "Europe/Warsaw")
    closure = tl["e_in"] - tl["e_out"] - tl["e_aux"] - tl["losses_pcs"] - tl["losses_battery"] - tl["d_stored"]
    assert closure.abs().max() < 1e-6
    both = res["balance"].join(tl.add_suffix("_true"))
    assert ((both["losses_pcs"] - both["losses_pcs_true"]).abs() / both["losses_pcs_true"]).max() < 0.005
    assert abs(both["losses_battery"].sum() / both["losses_battery_true"].sum() - 1) < 0.005


def test_forecast_brackets_truth(demo):
    sim, frame, res = demo
    fc = res["forecast"]
    plant = Plant()
    last = sim.truth["tests"][-1]
    yrs = fc["curve"]["years"].to_numpy()
    n = last["efc_total"] + fc["efc_per_year"] * np.maximum(yrs - last["years"], 0)
    e = plant.e0_stored_mwh * (1 - plant.fade_a * np.sqrt(yrs) - plant.fade_b * n / 1000)
    aux = last["e_use_mwh"] * plant.eta_pcs * plant.eta_batt - last["e_meter_mwh"]
    soh = (e * plant.eta_pcs * plant.eta_batt - aux) / plant.energy_mwh
    sel = (yrs >= 3) & (yrs <= 8)
    inside = (soh[sel] >= fc["curve"]["soh_p10"].to_numpy()[sel] - 0.003) & \
             (soh[sel] <= fc["curve"]["soh_p90"].to_numpy()[sel] + 0.003)
    assert inside.all()
    assert res["handoff"]["below_requirement"]["central"] is not None


def test_reports(demo, tmp_path):
    sim, frame, res = demo
    x = write_excel(res, "2026-08", tmp_path / "r.xlsx")
    h = write_html(res, "2026-08", tmp_path / "r.html")
    wb = load_workbook(x)
    assert {"KPI", "Monthly", "Daily", "Violations", "LD (illustrative)", "Capacity tests"} <= set(wb.sheetnames)
    assert wb["KPI"].conditional_formatting and wb["Monthly"].conditional_formatting
    text = h.read_text(encoding="utf-8")
    assert "illustrative" in text and "2026-08" in text and "<img" in text


def test_cli_run_on_a_short_export(demo, tmp_path):
    sim, frame, res = demo
    long = sim.long
    ts = pd.to_datetime(long["timestamp"])
    short = long[(ts >= "2025-10-20") & (ts < "2025-11-05")]       # includes the 25-hour day
    path = tmp_path / "ems.csv"
    short.to_csv(path, index=False)
    ev = sim.events
    ev.to_csv(tmp_path / "events.csv", index=False)
    from bess_ops.study import example_path
    rc = main(["run", str(path), "--mapping", str(example_path("mapping.example.yaml")),
               "--site", str(example_path("site.example.yaml")), "--warranty",
               str(example_path("warranty.example.yaml")), "--ltsa", str(example_path("ltsa.example.yaml")),
               "--events", str(tmp_path / "events.csv"), "--month", "2025-10", "--out", str(tmp_path / "out")])
    assert rc == 0
    assert (tmp_path / "out" / "report_2025-10.xlsx").exists()
    daily = pd.read_csv(tmp_path / "out" / "rte_daily.csv", index_col=0)
    assert daily.loc["2025-10-26", "hours"] == 25
